"""Общий код parser.py и parser2.py: сеть, конфиг, прогресс, обработка, XLSX."""
import html as html_module
import json
import logging
import os
import re
import time
from datetime import datetime
from urllib.parse import urlparse, parse_qs

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_SESSION = None


def get_session():
    global _SESSION
    if _SESSION is None:
        _SESSION = requests.Session()
    return _SESSION


def session_get(url, logger, verbose, retries=3, **kwargs):
    """GET с переиспользованием соединения и ретраями (429/5xx/сеть)."""
    delays = (1, 2, 4)
    for attempt in range(retries + 1):
        try:
            r = get_session().get(url, timeout=15, **kwargs)
            if r.status_code == 429:
                wait = int(r.headers.get('Retry-After', delays[min(attempt, 2)]))
                if verbose:
                    logger.warning(f"⏳ 429, ждём {wait}с (попытка {attempt+1})")
                time.sleep(wait)
                continue
            r.raise_for_status()
            return r
        except Exception as e:
            if attempt == retries:
                raise
            if verbose:
                logger.warning(f"⚠️ Повтор {attempt+1}/{retries}: {e}")
            time.sleep(delays[min(attempt, 2)])


def setup_logging(logger_name, log_file, verbose=False):
    logger = logging.getLogger(logger_name)
    logger.handlers.clear()

    if verbose:
        logger.setLevel(logging.DEBUG)
        file_handler = logging.FileHandler(log_file, encoding='utf-8', mode='w')
        file_handler.setLevel(logging.DEBUG)
        formatter = logging.Formatter('%(asctime)s - %(levelname)s - %(message)s',
                                      datefmt='%Y-%m-%d %H:%M:%S')
        file_handler.setFormatter(formatter)
        logger.addHandler(file_handler)
        print(f"📝 Логирование ВКЛ | Логи пишутся в: {log_file}")
    else:
        logger.setLevel(logging.CRITICAL + 1)
        logger.addHandler(logging.NullHandler())
        print(f"🔇 Логирование ВЫКЛ | Файл {log_file} не создаётся")

    return logger


def load_config(config_file='config.json', sort_default=False):
    default = {
        "api_token": "ВАШ_ТОКЕН_СЮДА",
        "max_vacancies": 100,
        "start_page": 0,
        "end_page": -1,
        "per_page": 100,
        "verbose": False,
        "sort_by_date": sort_default,
    }
    if os.path.exists(config_file):
        try:
            with open(config_file, 'r', encoding='utf-8') as f:
                cfg = json.load(f)
            for k, v in default.items():
                if k not in cfg:
                    cfg[k] = v
            return cfg
        except Exception as e:
            print(f"⚠️ Ошибка чтения конфига: {e}")
            return default
    with open(config_file, 'w', encoding='utf-8') as f:
        json.dump(default, f, ensure_ascii=False, indent=4)
    print(f"📝 Создан {config_file}. Заполните api_token.")
    return default


def save_config(cfg, config_file='config.json'):
    with open(config_file, 'w', encoding='utf-8') as f:
        json.dump(cfg, f, ensure_ascii=False, indent=4)


def save_progress(progress_file, progress_data):
    try:
        with open(progress_file, 'w', encoding='utf-8') as f:
            json.dump(progress_data, f, ensure_ascii=False, indent=2)
    except Exception as e:
        print(f"⚠️ Ошибка сохранения прогресса: {e}")


def load_progress(progress_file):
    if not os.path.exists(progress_file):
        return None
    try:
        with open(progress_file, 'r', encoding='utf-8') as f:
            return json.load(f)
    except Exception as e:
        print(f"⚠️ Ошибка чтения прогресса: {e}")
        return None


def clear_progress(progress_file):
    if os.path.exists(progress_file):
        try:
            os.remove(progress_file)
            print("🗑️ Файл прогресса удалён")
        except Exception:
            pass


def parse_search_url(url):
    parsed = urlparse(url)
    params = parse_qs(parsed.query)

    IGNORE_PARAMS = {
        'ored_clusters', 'enable_snippets', 'hhtmFromLabel', 'hhtmFrom',
        'search_session_id', 'suggest', 'host', 'cluster', 'no_magic',
        'page', 'per_page', 'L_save_area'
    }

    return {k: (v[0] if len(v) == 1 else v) for k, v in params.items() if k not in IGNORE_PARAMS}


def html_to_text(html_str):
    if not html_str:
        return '-'
    text = re.sub(r'<[^>]+>', ' ', html_str)
    text = html_module.unescape(text)
    text = re.sub(r'\s+', ' ', text).strip()
    return text if text else '-'


def extract_duties(text):
    if not text or text == '-':
        return '-'
    patterns = [
        r'Обязанности[:\s]*([\s\S]*?)(?=(?:Требования|Условия|Ключевые навыки|Что тебя ждет|$))',
        r'Что предстоит делать[:\s]*([\s\S]*?)(?=(?:Требования|Условия|$))',
        r'Чем предстоит заниматься[:\s]*([\s\S]*?)(?=(?:Мы ожидаем|Мы предлагаем|$))'
    ]
    for p in patterns:
        m = re.search(p, text, re.IGNORECASE)
        if m and len(m.group(1)) > 50:
            return m.group(1).strip()
    return text[:500] if len(text) > 500 else text


def format_date(date_str):
    try:
        dt = datetime.fromisoformat(date_str.replace('Z', '+00:00'))
        months = ['января', 'февраля', 'марта', 'апреля', 'мая', 'июня',
                  'июля', 'августа', 'сентября', 'октября', 'ноября', 'декабря']
        return f"{dt.day:02d} {months[dt.month-1]} {dt.year}"
    except Exception:
        return '-'


def parse_date_str(date_str):
    if not date_str or date_str == '-':
        return datetime.min
    months = {'января': 1, 'февраля': 2, 'марта': 3, 'апреля': 4, 'мая': 5, 'июня': 6,
              'июля': 7, 'августа': 8, 'сентября': 9, 'октября': 10, 'ноября': 11, 'декабря': 12}
    try:
        parts = date_str.split()
        day = int(parts[0])
        month = months.get(parts[1], 1)
        year = int(parts[2])
        return datetime(year, month, day)
    except Exception:
        return datetime.min


def sort_vacancies_by_date(vacancies):
    return sorted(vacancies, key=lambda v: parse_date_str(v.get('date', '-')), reverse=True)


def process_vacancy_details(details, index):
    if not details:
        return None

    name = details.get('name', '-')
    desc = html_to_text(details.get('description', ''))
    duties = extract_duties(desc)

    skills = '\n'.join(s['name'] for s in details.get('key_skills', []) if s.get('name')) or '-'
    employer = details.get('employer', {}).get('name', '-') if details.get('employer') else '-'
    city = details.get('area', {}).get('name', '-') if details.get('area') else '-'
    date = format_date(details.get('published_at', ''))
    url_vac = details.get('alternate_url', f"https://hh.ru/vacancy/{details.get('id', '')}")

    sal = details.get('salary')
    if sal:
        sf, st, sc = sal.get('from', ''), sal.get('to', ''), sal.get('currency', '')
        salary = f"{sf} - {st} {sc}" if sf and st else f"от {sf} {sc}" if sf else f"до {st} {sc}" if st else '-'
    else:
        salary = '-'

    schedule_parts = []

    experience = details.get('experience', {}).get('name', '')
    if experience:
        schedule_parts.append(f"Опыт работы: {experience}")

    if details.get('internship'):
        schedule_parts.append("Стажировка")

    employment_form = details.get('employment_form', {}).get('name', '')
    if employment_form:
        schedule_parts.append(employment_form)

    labor_contract = details.get('accept_labor_contract', False)
    civil_law = details.get('civil_law_contracts', [])

    contract_types = []
    if labor_contract:
        contract_types.append("Трудовой договор")
    if civil_law:
        for contract in civil_law:
            if isinstance(contract, dict):
                contract_types.append(f"Договор ГПХ ({contract.get('name', '')})")
            else:
                contract_types.append("Договор ГПХ")

    if contract_types:
        schedule_parts.append(f"Оформление: {', '.join(contract_types)}")

    work_schedule = details.get('work_schedule_by_days', [])
    if work_schedule:
        schedule_names = []
        for ws in work_schedule:
            if isinstance(ws, dict):
                schedule_names.append(ws.get('name', ''))
            else:
                schedule_names.append(str(ws))
        if schedule_names:
            schedule_parts.append(f"График: {', '.join(schedule_names)}")

    working_hours = details.get('working_hours', [])
    if working_hours:
        hours_names = []
        for wh in working_hours:
            if isinstance(wh, dict):
                hours_names.append(wh.get('name', ''))
            else:
                hours_names.append(str(wh))
        if hours_names:
            schedule_parts.append(f"Рабочие часы: {', '.join(hours_names)}")

    work_formats = details.get('work_format', [])
    if work_formats:
        format_names = []
        for wf in work_formats:
            if isinstance(wf, dict):
                wf_name = wf.get('name', '')
                if wf_name == 'На месте работодателя':
                    format_names.append('на месте работодателя')
                elif wf_name == 'Удалённо':
                    format_names.append('удалённо')
                elif wf_name == 'Гибрид':
                    format_names.append('гибрид')
                else:
                    format_names.append(wf_name.lower())
            else:
                format_names.append(str(wf))

        if format_names:
            schedule_parts.append(f"Формат работы: {', '.join(format_names)}")

    if details.get('night_shifts'):
        schedule_parts.append("Ночные смены")

    schedule_text = '\n'.join(schedule_parts) if schedule_parts else '-'

    return {
        'num': index,
        'id': details.get('id', ''),
        'url': url_vac,
        'name': name,
        'date': date,
        'desc': desc,
        'duties': duties,
        'skills': skills,
        'employer': employer,
        'city': city,
        'schedule': schedule_text,
        'salary': salary
    }


def save_to_xlsx(vacancies, filename):
    print("💾 Создание XLSX файла...")

    wb = Workbook()
    ws = wb.active
    ws.title = 'Вакансии'

    headers = ['№', 'ID', 'Ссылка', 'Название', 'Дата публикации', 'Описание',
               'Обязанности', 'Навыки', 'Работодатель', 'Город', 'График', 'Зарплата']

    ws.append(headers)

    header_font = Font(bold=True, color='000000', size=11)
    header_fill = PatternFill(start_color='0F8CFF', end_color='0F8CFF', fill_type='solid')
    header_align = Alignment(horizontal='left', vertical='center', wrap_text=True)

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    ws.row_dimensions[1].height = 25

    cell_align = Alignment(vertical='top', wrap_text=True)
    link_font = Font(color='0563C1', underline='single')

    for vac in vacancies:
        row_idx = ws.max_row + 1

        ws.cell(row=row_idx, column=1, value=vac['num']).alignment = cell_align
        ws.cell(row=row_idx, column=2, value=vac['id']).alignment = cell_align

        link_cell = ws.cell(row=row_idx, column=3, value=vac['url'])
        link_cell.hyperlink = vac['url']
        link_cell.font = link_font
        link_cell.alignment = cell_align

        ws.cell(row=row_idx, column=4, value=vac['name']).alignment = cell_align
        ws.cell(row=row_idx, column=5, value=vac['date']).alignment = cell_align
        ws.cell(row=row_idx, column=6, value=vac['desc']).alignment = cell_align
        ws.cell(row=row_idx, column=7, value=vac['duties']).alignment = cell_align
        ws.cell(row=row_idx, column=8, value=vac['skills']).alignment = cell_align
        ws.cell(row=row_idx, column=9, value=vac['employer']).alignment = cell_align
        ws.cell(row=row_idx, column=10, value=vac['city']).alignment = cell_align
        ws.cell(row=row_idx, column=11, value=vac['schedule']).alignment = cell_align
        ws.cell(row=row_idx, column=12, value=vac['salary']).alignment = cell_align

        schedule_text = str(vac['schedule'])
        if '\n' in schedule_text:
            line_count = schedule_text.count('\n') + 1
            ws.row_dimensions[row_idx].height = max(15, line_count * 15)

    for col_idx in range(1, 13):
        ws.column_dimensions[get_column_letter(col_idx)].width = 25

    wb.save(filename)
    print(f"✅ XLSX сохранён: {os.path.abspath(filename)}")
    print(f"📊 Строк всего: {ws.max_row}")
