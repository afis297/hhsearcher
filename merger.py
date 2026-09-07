import os
import glob
from tkinter import Tk, filedialog, messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from datetime import datetime

# ============================================================
# 1. ЧТЕНИЕ XLSX
# ============================================================
def read_xlsx(filepath):
    """Читает .xlsx, возвращает (headers, rows)."""
    try:
        wb = load_workbook(filepath, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return None, None
        headers = [str(h) if h is not None else '' for h in rows[0]]
        data = [[c if c is not None else '' for c in r[:len(headers)]] for r in rows[1:]]
        for r in data:
            while len(r) < len(headers):
                r.append('')
        return headers, data
    except Exception as e:
        print(f"  ⚠️ Ошибка чтения {os.path.basename(filepath)}: {e}")
        return None, None

# ============================================================
# 2. ВЫБОР ФАЙЛОВ
# ============================================================
def select_files_interactive(files):
    """Интерактивный выбор файлов по номерам"""
    if not files:
        return []

    print(f"\n📂 Доступные файлы XLSX:")
    print("-" * 50)
    for idx, fpath in enumerate(files, 1):
        size = os.path.getsize(fpath) / 1024
        print(f"  {idx:2}. {os.path.basename(fpath):<30} ({size:.1f} КБ)")
    print("-" * 50)
    print("  0 - Выбрать все файлы")

    while True:
        choice = input("Введите номера (через запятую) или 0: ").strip()
        if not choice:
            continue

        if choice == '0':
            print(f"✅ Выбрано все файлы: {len(files)}")
            return files

        try:
            indices = [int(x.strip()) for x in choice.split(',')]
            selected = [files[i-1] for i in indices if 1 <= i <= len(files)]

            if selected:
                print(f"✅ Выбрано файлов: {len(selected)}")
                return selected
            else:
                print("❌ Номера вне диапазона. Попробуйте снова.")
        except ValueError:
            print(" Введите числа через запятую (например: 1,3,5)")

# ============================================================
# 3. СОХРАНЕНИЕ
# ============================================================
def save_as_xlsx(headers, rows, filepath):
    """Сохраняет headers+rows в .xlsx с заголовками и гиперссылками"""
    print("   📝 Формирование XLSX...")
    wb = Workbook()
    ws = wb.active
    ws.title = 'Вакансии'

    header_font = Font(bold=True, color='000000', size=11)
    header_fill = PatternFill(start_color='0F8CFF', end_color='0F8CFF', fill_type='solid')
    header_align = Alignment(horizontal='left', vertical='center', wrap_text=True)
    cell_align = Alignment(vertical='top', wrap_text=True)
    link_font = Font(color='0563C1', underline='single')

    link_col_idx = headers.index('Ссылка') if 'Ссылка' in headers else None

    for col_idx, col_name in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    for row_data in rows:
        row_idx = ws.max_row + 1
        for col_idx, val in enumerate(row_data, 1):
            cell_value = str(val) if val != '' else ''
            cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
            cell.alignment = cell_align
            if link_col_idx is not None and col_idx == link_col_idx + 1 and cell_value.startswith('http'):
                cell.hyperlink = cell_value
                cell.font = link_font

    for col_idx in range(1, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = 25

    wb.save(filepath)
    print(f"   ✅ XLSX сохранён: {os.path.basename(filepath)}")

# ============================================================
# 4. ГЛАВНАЯ ЛОГИКА
# ============================================================
def main():
    root = Tk()
    root.withdraw()

    print("\n" + "="*60)
    print("  🔄 ОБЪЕДИНИТЕЛЬ ФАЙЛОВ ВАКАНСИЙ (v2.0)")
    print("="*60 + "\n")

    xlsx_files = sorted(glob.glob('*.xlsx'))

    if not xlsx_files:
        print("❌ В папке нет файлов .xlsx")
        messagebox.showwarning("Внимание", "Файлы для объединения не найдены")
        root.destroy()
        return

    print(f"📂 Найдено XLSX: {len(xlsx_files)}")

    selected = select_files_interactive(xlsx_files)

    if not selected:
        print("❌ Файлы не выбраны. Выход.")
        root.destroy()
        return

    print(f"\n Всего выбрано: {len(selected)} файл(ов)")

    master_cols = None
    merged = []

    for i, fpath in enumerate(selected, 1):
        print(f"\n[{i}/{len(selected)}] Чтение: {os.path.basename(fpath)}")

        headers, rows = read_xlsx(fpath)

        if not headers or not rows:
            print("   ⚠️ Пропуск (пустой или ошибка)")
            continue

        print(f"   📊 Строк: {len(rows)} | Колонки: {headers}")

        if master_cols is None:
            master_cols = headers
            print("   ✅ Установлены эталонные колонки")
            merged.extend(rows)
        else:
            idx = {h: n for n, h in enumerate(headers)}
            merged.extend([[r[idx[c]] if c in idx else '' for c in master_cols] for r in rows])
            print("   ✅ Выровнено по эталону")

    if not merged:
        print("\n❌ Нет данных для объединения")
        root.destroy()
        return

    print("\n🔄 ОБЪЕДИНЕНИЕ...")
    print(f"✅ Строк до очистки: {len(merged)}")

    if 'ID' in master_cols:
        id_col = master_cols.index('ID')
        seen, deduped = set(), []
        for r in merged:
            if r[id_col] not in seen:
                seen.add(r[id_col])
                deduped.append(r)
        if len(deduped) < len(merged):
            print(f"️ Удалено дубликатов: {len(merged) - len(deduped)}")
        merged = deduped

    if '№' in master_cols:
        num_col = master_cols.index('№')
        for n, r in enumerate(merged, 1):
            r[num_col] = n

    print(f"✅ Итого строк: {len(merged)}")

    timestamp = datetime.now().strftime('%Y-%m-%d_%H-%M')
    path = filedialog.asksaveasfilename(
        title="Сохранить XLSX",
        defaultextension=".xlsx",
        filetypes=[("Excel", "*.xlsx")],
        initialfile=f"hh_merged_{timestamp}.xlsx"
    )

    if path:
        save_as_xlsx(master_cols, merged, path)
        print("\n" + "="*60)
        print("🎉 ГОТОВО! Файл сохранён.")
        print("="*60)
        messagebox.showinfo("Успех", f"Файл сохранён:\n{path}")

    root.destroy()

if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\n⏹ Прервано пользователем")
    except Exception as e:
        print(f"\n❌ Критическая ошибка: {e}")
        import traceback
        traceback.print_exc()
        input("\nНажмите Enter для выхода...")
