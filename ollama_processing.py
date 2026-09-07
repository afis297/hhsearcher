import json
import time
import sys
import subprocess
import requests
from pathlib import Path
from openpyxl import Workbook, load_workbook
from tqdm import tqdm
from tkinter import Tk, filedialog

# ---------------------- НАСТРОЙКИ ----------------------
OLLAMA_URL = "http://localhost:11434/api/generate"
PROGRESS_EXT = ".progress.json"
# -------------------------------------------------------

# ====================== УПРАВЛЕНИЕ СЕРВЕРОМ OLLAMA ======================
class OllamaManager:
    """Запускает и останавливает сервер ollama serve при входе/выходе из контекста."""
    def __init__(self):
        self.process = None
        self._started_by_us = False

    def _is_ollama_running(self) -> bool:
        """Проверяет доступность API Ollama."""
        try:
            r = requests.get("http://localhost:11434/api/tags", timeout=2)
            return r.status_code == 200
        except Exception:
            return False

    def start(self):
        """Запускает сервер, если он ещё не запущен."""
        if self._is_ollama_running():
            print("✅ Ollama уже запущен (используем существующий экземпляр).")
            self._started_by_us = False
            return

        print("🚀 Запуск Ollama сервера...")
        # Для Windows скрываем окно консоли, для Unix – стандартный запуск
        creation_flags = subprocess.CREATE_NO_WINDOW if sys.platform == 'win32' else 0
        self.process = subprocess.Popen(
            ['ollama', 'serve'],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=creation_flags
        )
        self._started_by_us = True

        # Ждём готовности (до 15 секунд)
        for _ in range(15):
            time.sleep(1)
            if self._is_ollama_running():
                print("✅ Ollama сервер готов.")
                return
        raise RuntimeError("Не удалось запустить Ollama сервер за 15 секунд.")

    def stop(self):
        """Останавливает сервер, только если он был запущен этим менеджером."""
        if self._started_by_us and self.process and self.process.poll() is None:
            print("🛑 Останавливаем Ollama сервер...")
            self.process.terminate()
            try:
                self.process.wait(timeout=5)
            except subprocess.TimeoutExpired:
                self.process.kill()
            print("✅ Ollama остановлен.")
        elif self._started_by_us:
            print("⚠️ Ollama уже завершился неожиданно.")
        else:
            print("ℹ️ Ollama не был запущен скриптом, оставляем как есть.")

    def __enter__(self):
        self.start()
        return self

    def __exit__(self, exc_type, exc_val, exc_tb):
        self.stop()
        # Не подавляем исключения
        return False

# ---------------------- ОСТАЛЬНЫЕ ФУНКЦИИ (БЕЗ ИЗМЕНЕНИЙ) ----------------------
def select_model():
    print("\nДоступные модели Ollama:")
    print("1. qwen2.5:1.5b (быстро, меньше памяти)")
    print("2. qwen2.5:7b   (медленнее, качественнее)")
    choice = input("Выберите модель (1 или 2): ").strip()
    if choice == "1":
        return "qwen2.5:1.5b"
    elif choice == "2":
        return "qwen2.5:7b"
    else:
        print("Неверный выбор, использую qwen2.5:7b")
        return "qwen2.5:7b"

def select_file(prompt="Выберите Excel-файл с вакансиями"):
    try:
        Tk().withdraw()
        path = filedialog.askopenfilename(title=prompt, filetypes=[("Excel files", "*.xlsx")])
        if path:
            return Path(path)
    except Exception:
        pass
    path = input(f"{prompt} (или укажите путь): ").strip()
    return Path(path)

def select_output_file(default_name="vacancies_with_duties.xlsx"):
    out = input(f"Имя выходного файла (Enter = {default_name}): ").strip()
    if not out:
        out = default_name
    return out

def load_progress(progress_file: Path) -> dict:
    if progress_file.exists():
        try:
            with open(progress_file, "r", encoding="utf-8") as f:
                data = json.load(f)
                return {int(k): v for k, v in data.items()}
        except Exception:
            return {}
    return {}

def save_progress(progress_file: Path, progress: dict):
    with open(progress_file, "w", encoding="utf-8") as f:
        json.dump(progress, f, ensure_ascii=False, indent=2)

def ask_continue(progress_file: Path) -> bool:
    if not progress_file.exists():
        return True
    answer = input("Найден сохранённый прогресс. Продолжить с места остановки? (y/n): ").strip().lower()
    return answer == 'y'

def extract_duties_sync(description: str, model: str) -> tuple[str, float]:
    """Синхронный вызов Ollama (через requests)."""
    if not isinstance(description, str) or len(description.strip()) < 50:
        return "", 0.0
    prompt = f"""
Из текста вакансии извлеки **только обязанности** (что нужно делать сотруднику).  
Игнорируй требования, условия, бонусы, описание компании.  
Выдай в виде маркированного списка, каждый пункт начинается с символа "•".  
Если обязанностей нет, напиши "Не указаны".

Текст вакансии:
{description[:3000]}
"""
    payload = {
        "model": model,
        "prompt": prompt,
        "stream": False,
        "options": {"temperature": 0.1, "num_predict": 800}
    }
    start = time.perf_counter()
    try:
        resp = requests.post(OLLAMA_URL, json=payload, timeout=60)
        if resp.status_code == 200:
            result = resp.json().get("response", "").strip()
        else:
            result = ""
    except Exception as e:
        print(f"Ошибка: {e}")
        result = ""
    elapsed = time.perf_counter() - start
    return result, elapsed

def read_table(path: Path):
    """Читает первый лист xlsx, возвращает (headers, rows)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = [[c if c is not None else '' for c in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return [str(h) for h in rows[0]], rows[1:]


def write_table(headers, rows, path: str):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Вакансии'
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def process_sync(headers, rows, model: str, output_path: str, progress_file: Path):
    progress = load_progress(progress_file)
    total = len(rows)
    desc_col = headers.index('Описание') if 'Описание' in headers else None
    if 'Обязанности_ollama' not in headers:
        headers.append('Обязанности_ollama')
        for r in rows:
            r.append('')
    duty_col = headers.index('Обязанности_ollama')
    for i in range(total):
        if i not in progress and rows[i][duty_col]:
            progress[i] = rows[i][duty_col]
    pending_indices = [i for i in range(total) if i not in progress]
    if not pending_indices:
        write_table(headers, rows, output_path)
        print(f"✅ Все уже обработано. Результат сохранён в {output_path}")
        return

    times = []
    with tqdm(total=len(pending_indices), desc="Обработка вакансий", unit="шт") as pbar:
        for idx in pending_indices:
            description = rows[idx][desc_col] if desc_col is not None else ''
            result, elapsed = extract_duties_sync(str(description or ''), model)
            rows[idx][duty_col] = result
            progress[idx] = result
            times.append(elapsed)
            pbar.update(1)
            if len(progress) % 5 == 0:
                save_progress(progress_file, progress)
            if len(times) % 5 == 0:
                avg = sum(times) / len(times)
                pbar.set_postfix({"среднее": f"{avg:.2f}с"})

    save_progress(progress_file, progress)
    if times:
        avg_time = sum(times) / len(times)
        total_time = sum(times)
        print("\n=== Статистика обработки ===")
        print(f"Всего обработано (сессия): {len(times)}")
        print(f"Общее время: {total_time:.2f} сек ({total_time/60:.2f} мин)")
        print(f"Среднее время: {avg_time:.2f} сек")
        print(f"Макс: {max(times):.2f} сек, мин: {min(times):.2f} сек")

    write_table(headers, rows, output_path)
    print(f"\n✅ Результат сохранён в {output_path}")
    if progress_file.exists():
        progress_file.unlink()

# ---------------------- ТОЧКА ВХОДА С УПРАВЛЕНИЕМ OLLAMA ----------------------
def main():
    """Точка входа, вся обработка внутри менеджера Ollama."""
    with OllamaManager():
        print("=== Извлечение обязанностей из вакансий ===\n")
        model = select_model()
        input_file = select_file()
        if not input_file.exists():
            print(f"Файл {input_file} не найден.")
            return

        output_file = select_output_file()
        progress_file = input_file.with_suffix(PROGRESS_EXT)

        if not ask_continue(progress_file):
            if progress_file.exists():
                progress_file.unlink()

        print(f"\nФайл: {input_file}")
        print(f"Модель: {model}")
        print("Начинаем обработку...\n")

        headers, rows = read_table(input_file)
        print(f"Загружено {len(rows)} вакансий.")

        process_sync(headers, rows, model, output_file, progress_file)

if __name__ == "__main__":
    main()